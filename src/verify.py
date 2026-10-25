"""Independently check the five KPIs using workbook rows and Decimal arithmetic."""
from collections import defaultdict
from decimal import Decimal
from pathlib import Path
import csv
import json
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
book = load_workbook(ROOT/'data/raw/Online Retail.xlsx',read_only=True,data_only=True)
rows = book.active.iter_rows(values_only=True)
next(rows)
seen = set()
invoices = set()
customer_orders = defaultdict(set)
sales = Decimal(0)
raw_count = purchase_count = 0
for row in rows:
    raw_count += 1
    if row in seen:
        continue
    seen.add(row)
    invoice, stock, description, quantity, date, price, customer, country = row
    if str(invoice).upper().startswith('C') or quantity <= 0 or price <= 0:
        continue
    purchase_count += 1
    sales += Decimal(str(price)) * quantity
    invoices.add(str(invoice))
    if customer is not None:
        customer_orders[str(customer)].add(str(invoice))
book.close()
actual = {
    'gross_purchase_sales_gbp':float(sales),
    'purchase_orders':len(invoices),
    'identified_purchasing_customers':len(customer_orders),
    'average_purchase_order_value_gbp':float(sales/len(invoices)),
    'identified_repeat_customer_rate':sum(len(v)>=2 for v in customer_orders.values())/len(customer_orders),
}
expected = json.loads((ROOT/'results/kpis.json').read_text())
for key,value in actual.items():
    assert abs(value-expected[key]) < 1e-9, (key,value,expected[key])
quality = json.loads((ROOT/'results/data_quality.json').read_text())
assert raw_count == quality['source_rows']
assert purchase_count == quality['purchase_rows']
audit = list(csv.DictReader((ROOT/'results/cleaning_audit.csv').open()))
assert raw_count-sum(int(row['removed']) for row in audit) == purchase_count
report = {'method':'Independent openpyxl row iteration and Decimal arithmetic',
          'five_kpis_match':True,'source_and_purchase_counts_match':True,
          'absolute_numeric_tolerance':1e-9,'kpis':actual}
(ROOT/'results/independent_verification.json').write_text(json.dumps(report,indent=2)+'\n')
print(json.dumps(report,indent=2))
